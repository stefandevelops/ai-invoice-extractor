from __future__ import annotations

import csv
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from extractor.models import ReviewedDocument
from extractor.money import fmt_money

DOCUMENT_COLUMNS = [
    "status",
    "issues",
    "filename",
    "doc_type",
    "supplier",
    "document_number",
    "invoice_date",
    "due_date",
    "delivery_date",
    "po_number",
    "currency",
    "subtotal",
    "vat_amount",
    "total",
    "line_count",
    "parser",
]

LINE_COLUMNS = [
    "filename",
    "document_number",
    "description",
    "quantity",
    "unit_price",
    "vat_rate",
    "line_total",
]

FLAG_FILL = PatternFill("solid", fgColor="FDECEC")
CLEAN_FILL = PatternFill("solid", fgColor="EEF6F0")
WARN_FILL = PatternFill("solid", fgColor="FFF4D6")
HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def document_rows(reviewed: list[ReviewedDocument]) -> list[dict[str, str]]:
    rows = []
    for row in reviewed:
        doc = row.document
        rows.append(
            {
                "status": row.status_label,
                "issues": " | ".join(issue.message for issue in row.issues),
                "filename": doc.source_filename,
                "doc_type": doc.doc_type,
                "supplier": doc.supplier,
                "document_number": doc.document_number,
                "invoice_date": _iso(doc.invoice_date),
                "due_date": _iso(doc.due_date),
                "delivery_date": _iso(doc.delivery_date),
                "po_number": doc.po_number or "",
                "currency": doc.currency,
                "subtotal": _amount(doc.subtotal),
                "vat_amount": _amount(doc.vat_amount),
                "total": _amount(doc.total),
                "line_count": str(len(doc.line_items)),
                "parser": doc.parser,
            }
        )
    return rows


def line_rows(reviewed: list[ReviewedDocument]) -> list[dict[str, str]]:
    rows = []
    for row in reviewed:
        doc = row.document
        for item in doc.line_items:
            rows.append(
                {
                    "filename": doc.source_filename,
                    "document_number": doc.document_number,
                    "description": item.description,
                    "quantity": _amount(item.quantity),
                    "unit_price": _amount(item.unit_price),
                    "vat_rate": str(item.vat_rate),
                    "line_total": _amount(item.line_total),
                }
            )
    return rows


def to_csv(reviewed: list[ReviewedDocument]) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=DOCUMENT_COLUMNS)
    writer.writeheader()
    writer.writerows(document_rows(reviewed))
    return buffer.getvalue().encode("utf-8")


def to_xlsx(reviewed: list[ReviewedDocument]) -> bytes:
    workbook = Workbook()
    documents_sheet = workbook.active
    documents_sheet.title = "Documents"
    _write_sheet(documents_sheet, DOCUMENT_COLUMNS, document_rows(reviewed), flag_status=True)

    lines_sheet = workbook.create_sheet("Line items")
    _write_sheet(lines_sheet, LINE_COLUMNS, line_rows(reviewed), flag_status=False)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def review_table_records(reviewed: list[ReviewedDocument]) -> list[dict[str, str]]:
    records = []
    for row in reviewed:
        doc = row.document
        records.append(
            {
                "Status": row.status_label,
                "Supplier": doc.supplier,
                "Number": doc.document_number,
                "Type": "Invoice" if doc.doc_type == "invoice" else "Delivery note",
                "Date": _iso(doc.invoice_date or doc.delivery_date),
                "PO": doc.po_number or "—",
                "Total": fmt_money(doc.total, doc.currency),
                "Issues": "; ".join(issue.message for issue in row.issues) or "—",
                "File": doc.source_filename,
            }
        )
    return records


def _write_sheet(sheet, columns: list[str], rows: list[dict[str, str]], flag_status: bool) -> None:
    sheet.append(columns)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
        current = sheet[sheet.max_row]
        fill = None
        if flag_status:
            status = row.get("status", "")
            if status == "Flagged":
                fill = FLAG_FILL
            elif status == "Warning":
                fill = WARN_FILL
            else:
                fill = CLEAN_FILL
        for cell in current:
            if fill is not None:
                cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {column: len(column) + 2 for column in columns}
    for row in rows:
        for column in columns:
            widths[column] = min(max(widths[column], len(row.get(column, "")) + 2), 60)
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths[column]
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"


def _iso(value) -> str:
    return value.isoformat() if value else ""


def _amount(value: Decimal | None) -> str:
    return f"{value:.2f}" if value is not None else ""
