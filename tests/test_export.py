from decimal import Decimal

from extractor.export import document_rows, to_csv, to_xlsx
from extractor.models import ExtractedDocument, Issue, LineItem, ReviewedDocument


def _reviewed() -> list[ReviewedDocument]:
    doc = ExtractedDocument(
        source_filename="a.pdf",
        doc_type="invoice",
        supplier="Nordik Office Supplies AB",
        document_number="INV-1",
        po_number="PO-1",
        currency="EUR",
        line_items=[
            LineItem(
                description="Paper",
                quantity=Decimal("2"),
                unit_price=Decimal("10.00"),
                vat_rate=Decimal("0.25"),
                line_total=Decimal("20.00"),
            )
        ],
        subtotal=Decimal("20.00"),
        vat_amount=Decimal("5.00"),
        total=Decimal("25.00"),
    )
    return [
        ReviewedDocument(
            document=doc,
            issues=[Issue(code="vat_mismatch", severity="error", message="VAT looks wrong.")],
        )
    ]


def test_csv_includes_status_and_reason():
    payload = to_csv(_reviewed()).decode("utf-8")
    assert "Flagged" in payload
    assert "VAT looks wrong." in payload
    assert "INV-1" in payload


def test_xlsx_has_two_sheets():
    from io import BytesIO

    from openpyxl import load_workbook

    book = load_workbook(BytesIO(to_xlsx(_reviewed())))
    assert book.sheetnames == ["Documents", "Line items"]
    assert book["Documents"]["A2"].value == "Flagged"
    assert book["Line items"]["C2"].value == "Paper"


def test_document_rows_flatten_issues():
    rows = document_rows(_reviewed())
    assert rows[0]["status"] == "Flagged"
    assert rows[0]["line_count"] == "1"
